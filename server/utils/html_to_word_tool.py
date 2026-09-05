from docx import Document
from docx.shared import Pt, RGBColor, Inches, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.style import WD_STYLE_TYPE
from bs4 import BeautifulSoup, Tag
import re
import base64
from typing import Dict, List, Optional
import os

class HTMLToWordConverter:
    """HTML转Word转换器（企业级）"""
    
    def __init__(self, reference_docx=None):
        """
        初始化转换器
        
        Args:
            reference_docx: 参考文档模板路径
        """
        self.doc = Document(reference_docx) if reference_docx else Document()
        self.styles = self.doc.styles
        self.image_counter = 0
        self._define_custom_styles()
        self.tags = []
    
    def _define_custom_styles(self):
        """定义自定义样式"""
        # 标题样式
        for i in range(1, 6):
            style_name = f"Heading {i}"
            if style_name not in self.styles:
                style = self.styles.add_style(style_name, WD_STYLE_TYPE.PARAGRAPH)
                style.font.size = Pt(24 - (i-1)*2)
                style.font.bold = True
        
        # 代码样式
        if 'Code' not in self.styles:
            code_style = self.styles.add_style('Code', WD_STYLE_TYPE.PARAGRAPH)
            code_style.font.name = 'Courier New'
            code_style.font.size = Pt(10)
            code_style.paragraph_format.left_indent = Pt(20)
            code_style.paragraph_format.space_before = Pt(6)
            code_style.paragraph_format.space_after = Pt(6)
    
    def convert(self, html_content: str, output_path: str) -> str:
        """
        转换HTML到Word
        
        Args:
            html_content: HTML字符串
            output_path: 输出文件路径
            
        Returns:
            输出文件路径
        """
        # 清理HTML
        html_content = self._clean_html(html_content)
        
        # 解析HTML
        soup = BeautifulSoup(html_content, 'html.parser')
        
        # 转换主体内容
        body = soup.body if soup.body else soup
        self._convert_element(body)
        # 保存文档
        self.doc.save(output_path)
        return output_path
    
    def _clean_html(self, html: str) -> str:
        """清理HTML"""
        # 移除不需要的标签
        html = re.sub(r'<script.*?>.*?</script>', '', html, flags=re.DOTALL)
        html = re.sub(r'<style.*?>.*?</style>', '', html, flags=re.DOTALL)
        
        # 修复不完整的标签
        html = html.replace('<br>', '<br/>').replace('<hr>', '<hr/>')
        
        return html
    
    def _convert_element(self, element):
        """递归转换HTML元素"""
        if isinstance(element, str):
            self._add_text(element)
        elif isinstance(element, Tag):
            tag_name = element.name.lower()
            if tag_name not in self.tags:
                self.tags.append(tag_name)
            # 处理不同标签
            handler_name = f'_handle_{tag_name}'
            if hasattr(self, handler_name):
                handler = getattr(self, handler_name)
                handler(element)
            else:
                # 默认处理：递归处理子元素
                for child in element.children:
                    self._convert_element(child)
    
    def _add_text(self, text: str, style=None):
        """添加文本"""
        if not text.strip() and text != ' ':
            return
        
        # 获取当前段落
        if not self.doc.paragraphs or self._current_paragraph_is_empty():
            p = self.doc.add_paragraph()
        else:
            p = self.doc.paragraphs[-1]
        
        # 添加文本
        run = p.add_run(text)
        
        # 应用样式
        if style:
            for key, value in style.items():
                if hasattr(run.font, key):
                    setattr(run.font, key, value)
    
    def _handle_h1(self, element):
        self.doc.add_heading(element.get_text(strip=True), 0)
    
    def _handle_h2(self, element):
        self.doc.add_heading(element.get_text(strip=True), 1)
    
    def _handle_h3(self, element):
        self.doc.add_heading(element.get_text(strip=True), 2)
    
    def _handle_p(self, element):
        p = self.doc.add_paragraph()
        self._process_element_content(element, p)
    
    def _process_element_content(self, element, paragraph):
        """处理元素内容"""
        for child in element.children:
            if isinstance(child, str):
                paragraph.add_run(child)
            elif isinstance(child, Tag):
                if child.name in ['strong', 'b']:
                    self._process_bold_tag(child, paragraph)
                elif child.name in ['em', 'i']:
                    self._process_italic_tag(child, paragraph)
                elif child.name == 'u':
                    self._process_underline_tag(child, paragraph)
                elif child.name == 'span':
                    self._process_span_tag(child, paragraph)
                elif child.name == 'br':
                    paragraph.add_run('\n')
                elif child.name == 'img':  # 添加对图片的处理
                    self._handle_img_in_paragraph(child, paragraph)
                else:
                    # 递归处理其他标签
                    self._process_element_content(child, paragraph)

    def _handle_img_in_paragraph(self, element, paragraph):
      """在段落中处理图片"""
      src = element.get('src', '')
      if not src:
          return
      
      # 处理base64图片
      if src.startswith('data:image'):
          try:
              # 提取base64数据
              header, data = src.split(',', 1)
              img_data = base64.b64decode(data)
              
              # 创建临时文件
              import tempfile
              with tempfile.NamedTemporaryFile(delete=False, suffix='.png') as f:
                  f.write(img_data)
                  temp_path = f.name
              
              # 添加图片到当前段落
              # 确保段落是居中对齐
              paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
              
              # 获取图片尺寸
              width = Inches(3)  # 默认宽度3英寸
              if element.get('width'):
                  try:
                      if 'px' in element.get('width'):
                          px_width = int(element.get('width').replace('px', ''))
                          width = Pt(px_width)
                  except:
                      pass
              
              # 在段落中添加图片
              run = paragraph.add_run()
              run.add_picture(temp_path, width=width)
              
              # 清理临时文件
              import os
              os.unlink(temp_path)
              
          except Exception as e:
              print(f"段落内图片处理失败: {e}")


    def _process_bold_tag(self, element, paragraph):
        """处理加粗标签"""
        for child in element.children:
            if isinstance(child, str):
                run = paragraph.add_run(child)
                run.bold = True
            elif isinstance(child, Tag):
                if child.name in ['em', 'i']:
                    self._process_bold_italic_tag(child, paragraph)
                elif child.name == 'span':
                    self._process_span_with_formatting(child, paragraph, bold=True)
                else:
                    run = paragraph.add_run(child.get_text())
                    run.bold = True
    
    def _process_italic_tag(self, element, paragraph):
        """处理斜体标签"""
        for child in element.children:
            if isinstance(child, str):
                run = paragraph.add_run(child)
                run.italic = True
            elif isinstance(child, Tag):
                if child.name in ['strong', 'b']:
                    self._process_bold_italic_tag(child, paragraph, italic_first=True)
                elif child.name == 'span':
                    self._process_span_with_formatting(child, paragraph, italic=True)
                else:
                    run = paragraph.add_run(child.get_text())
                    run.italic = True
    
    def _process_underline_tag(self, element, paragraph):
        """处理下划线标签"""
        for child in element.children:
            if isinstance(child, str):
                run = paragraph.add_run(child)
                run.underline = True
            elif isinstance(child, Tag):
                if child.name in ['strong', 'b']:
                    run = paragraph.add_run(child.get_text())
                    run.bold = True
                    run.underline = True
                elif child.name in ['em', 'i']:
                    run = paragraph.add_run(child.get_text())
                    run.italic = True
                    run.underline = True
                elif child.name == 'span':
                    self._process_span_with_formatting(child, paragraph, underline=True)
                else:
                    run = paragraph.add_run(child.get_text())
                    run.underline = True
    
    def _process_bold_italic_tag(self, element, paragraph, italic_first=False):
        """处理加粗斜体标签"""
        for child in element.children:
            if isinstance(child, str):
                run = paragraph.add_run(child)
                run.bold = True
                run.italic = True
            elif isinstance(child, Tag):
                if child.name == 'span':
                    self._process_span_with_formatting(child, paragraph, bold=True, italic=True)
                else:
                    run = paragraph.add_run(child.get_text())
                    run.bold = True
                    run.italic = True
    
    def _process_span_tag(self, element, paragraph):
        """处理span标签 - 修复WPS兼容性问题"""
        # 解析内联样式
        style_dict = self._parse_inline_styles_to_dict(element.get('style', ''))
        
        for child in element.children:
            if isinstance(child, str):
                run = paragraph.add_run(child)
                self._apply_run_styles_for_wps(run, style_dict)
            elif isinstance(child, Tag):
                if child.name in ['strong', 'b']:
                    style_dict['bold'] = True
                    self._process_formatted_span_child(child, paragraph, style_dict)
                elif child.name in ['em', 'i']:
                    style_dict['italic'] = True
                    self._process_formatted_span_child(child, paragraph, style_dict)
                elif child.name == 'u':
                    style_dict['underline'] = True
                    self._process_formatted_span_child(child, paragraph, style_dict)
                elif child.name == 'span':
                    # 处理嵌套的span
                    nested_style = style_dict.copy()
                    nested_style.update(self._parse_inline_styles_to_dict(child.get('style', '')))
                    self._process_nested_span_tag(child, paragraph, nested_style)
                else:
                    self._process_formatted_span_child(child, paragraph, style_dict)
    
    def _process_span_with_formatting(self, element, paragraph, bold=False, italic=False, underline=False):
        """处理带格式的span标签"""
        # 解析内联样式
        style_dict = self._parse_inline_styles_to_dict(element.get('style', ''))
        
        # 应用格式
        if bold:
            style_dict['bold'] = True
        if italic:
            style_dict['italic'] = True
        if underline:
            style_dict['underline'] = True
        
        for child in element.children:
            if isinstance(child, str):
                run = paragraph.add_run(child)
                self._apply_run_styles_for_wps(run, style_dict)
            elif isinstance(child, Tag):
                if child.name in ['strong', 'b']:
                    style_dict['bold'] = True
                elif child.name in ['em', 'i']:
                    style_dict['italic'] = True
                elif child.name == 'u':
                    style_dict['underline'] = True
                
                self._process_formatted_span_child(child, paragraph, style_dict)
    
    def _process_nested_span_tag(self, element, paragraph, style_dict):
        """处理嵌套的span标签"""
        for child in element.children:
            if isinstance(child, str):
                run = paragraph.add_run(child)
                self._apply_run_styles_for_wps(run, style_dict)
            elif isinstance(child, Tag):
                if child.name in ['strong', 'b']:
                    style_dict['bold'] = True
                elif child.name in ['em', 'i']:
                    style_dict['italic'] = True
                elif child.name == 'u':
                    style_dict['underline'] = True
                
                self._process_formatted_span_child(child, paragraph, style_dict)
    
    def _process_formatted_span_child(self, element, paragraph, style_dict):
        """处理span内的格式化子元素"""
        for child in element.children:
            if isinstance(child, str):
                run = paragraph.add_run(child)
                self._apply_run_styles_for_wps(run, style_dict)
            elif isinstance(child, Tag):
                if child.name in ['strong', 'b']:
                    style_dict['bold'] = True
                elif child.name in ['em', 'i']:
                    style_dict['italic'] = True
                elif child.name == 'u':
                    style_dict['underline'] = True
                elif child.name == 'span':
                    # 处理嵌套span
                    nested_style = style_dict.copy()
                    nested_style.update(self._parse_inline_styles_to_dict(child.get('style', '')))
                    self._process_nested_span_tag(child, paragraph, nested_style)
                else:
                    self._process_formatted_span_child(child, paragraph, style_dict)
    
    def _parse_inline_styles_to_dict(self, style_str: str) -> dict:
        """解析内联样式为字典"""
        styles = {}
        if not style_str:
            return styles
            
        for style in style_str.split(';'):
            style = style.strip()
            if ':' in style:
                key, value = style.split(':', 1)
                key = key.strip()
                value = value.strip()
                styles[key] = value
        return styles
    
    def _apply_run_styles_for_wps(self, run, style_dict: dict):
        """为WPS应用样式 - 修复版本"""
        for prop, value in style_dict.items():
            if prop == 'color':
                # WPS兼容的颜色设置
                rgb_color = self._parse_color_for_wps(value)
                if rgb_color:
                    # 使用正确的颜色设置方法
                    run.font.color.rgb = rgb_color
            elif prop == 'font-size':
                if 'px' in value:
                    try:
                        size_px = float(value.replace('px', '').strip())
                        size_pt = size_px * 0.75
                        run.font.size = Pt(size_pt)
                    except:
                        pass
            elif prop == 'font-weight':
                if value in ['bold', '700', '800', '900']:
                    run.bold = True
            elif prop == 'font-style':
                if value == 'italic':
                    run.italic = True
            elif prop == 'text-decoration':
                if 'underline' in value:
                    run.underline = True
            elif prop == 'bold' and value is True:
                run.bold = True
            elif prop == 'italic' and value is True:
                run.italic = True
            elif prop == 'underline' and value is True:
                run.underline = True
    
    def _parse_color_for_wps(self, color_str: str) -> Optional[RGBColor]:
        """为WPS解析颜色值"""
        if not color_str:
            return None
            
        color_str = color_str.strip().lower()
        
        # 常见颜色映射
        color_map = {
            'red': RGBColor(255, 0, 0),
            'green': RGBColor(0, 128, 0),
            'blue': RGBColor(0, 0, 255),
            'black': RGBColor(0, 0, 0),
            'white': RGBColor(255, 255, 255),
            'gray': RGBColor(128, 128, 128),
            'grey': RGBColor(128, 128, 128),
            'yellow': RGBColor(255, 255, 0),
            'orange': RGBColor(255, 165, 0),
            'purple': RGBColor(128, 0, 128),
            'brown': RGBColor(165, 42, 42),
            'cyan': RGBColor(0, 255, 255),
            'magenta': RGBColor(255, 0, 255),
        }
        
        if color_str in color_map:
            return color_map[color_str]
        
        # 处理十六进制颜色
        if color_str.startswith('#'):
            hex_color = color_str[1:].strip()
            if len(hex_color) == 3:
                # 扩展简写颜色 #abc -> #aabbcc
                hex_color = ''.join([c*2 for c in hex_color])
            
            if len(hex_color) == 6:
                try:
                    r = int(hex_color[0:2], 16)
                    g = int(hex_color[2:4], 16)
                    b = int(hex_color[4:6], 16)
                    return RGBColor(r, g, b)
                except:
                    return None
        
        # 处理rgb/rgba颜色
        rgb_match = re.match(r'rgba?\(\s*(\d+)\s*,\s*(\d+)\s*,\s*(\d+)', color_str)
        if rgb_match:
            try:
                r, g, b = map(int, rgb_match.groups())
                return RGBColor(r, g, b)
            except:
                return None
        
        return None
    
    def _handle_ul(self, element):
        for li in element.find_all('li', recursive=False):
            p = self.doc.add_paragraph(style='List Bullet')
            self._convert_list_item(p, li)
    
    def _handle_ol(self, element):
        for i, li in enumerate(element.find_all('li', recursive=False), 1):
            p = self.doc.add_paragraph(style='List Number')
            self._convert_list_item(p, li)
    
    def _handle_table(self, element):
        # 计算表格行列
        rows = element.find_all('tr')
        if not rows:
            return
        
        col_count = max(len(row.find_all(['td', 'th'])) for row in rows)
        table = self.doc.add_table(rows=len(rows), cols=col_count)
        
        for i, row in enumerate(rows):
            cells = row.find_all(['td', 'th'])
            for j, cell in enumerate(cells):
                if j < col_count:
                    # 创建段落并添加内容
                    paragraph = table.cell(i, j).paragraphs[0]
                    self._process_element_content(cell, paragraph)
                    
                    # 设置表头样式
                    if cell.name == 'th':
                        for run in paragraph.runs:
                            run.bold = True
        
        # 应用表格样式
        if element.has_attr('border'):
            table.style = 'Table Grid'
    
    def _handle_img(self, element):
        """处理图片"""
        src = element.get('src', '')
        if not src:
            return
        
        # 处理base64图片
        if src.startswith('data:image'):
            try:
                # 提取base64数据
                header, data = src.split(',', 1)
                img_data = base64.b64decode(data)
                
                # 创建临时文件
                import tempfile
                with tempfile.NamedTemporaryFile(delete=False, suffix='.png') as f:
                    import random
                    path = os.path.join('.', f'{random.randint(1, 10)}_{random.randint(1, 10)}.png')
                    print(path)
                    with open(path, mode='w', encoding='utf-8') as p:
                        p.write(img_data)
                    f.write(img_data)
                    temp_path = f.name
                
                # 添加图片到文档
                if self.doc.paragraphs:
                    p = self.doc.add_paragraph()
                else:
                    p = self.doc.paragraphs[-1]
            
                # 设置段落对齐方式为居中
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                
                p.add_run().add_picture(temp_path, width=Inches(3))
                
                # 清理临时文件
                os.unlink(temp_path)
                
            except Exception as e:
                print(f"图片处理失败: {e}")
    
    def _current_paragraph_is_empty(self):
        """检查当前段落是否为空"""
        if not self.doc.paragraphs:
            return True
        last_para = self.doc.paragraphs[-1]
        return not last_para.text.strip()
    
    def _convert_list_item(self, paragraph, li_element):
        """转换列表项"""
        for child in li_element.children:
            if isinstance(child, str):
                paragraph.add_run(child)
            elif isinstance(child, Tag):
                if child.name in ['strong', 'b']:
                    run = paragraph.add_run(child.get_text())
                    run.bold = True
                elif child.name in ['em', 'i']:
                    run = paragraph.add_run(child.get_text())
                    run.italic = True
                elif child.name == 'span':
                    self._process_span_tag(child, paragraph)
                else:
                    # 递归处理嵌套元素
                    self._convert_element(child)